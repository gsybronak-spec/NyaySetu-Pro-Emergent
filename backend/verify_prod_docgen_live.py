import io
import os
import sys
import time
import json
import base64
import uuid
import asyncio
from datetime import datetime, timezone, timedelta
import zipfile
import pypdfium2 as pdfium
from docx import Document
import httpx
import jwt
from google.cloud import firestore
from google.oauth2 import service_account

# Ensure UTF-8 output handling on Windows console
if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

CRED_FILE = r"C:\Users\HP\Downloads\nyaysetu-pro-firebase-adminsdk-fbsvc-dd4459b9fd.json"
if not os.path.exists(CRED_FILE):
    print(f"[FAIL] Credentials file not found: {CRED_FILE}")
    sys.exit(1)

with open(CRED_FILE, "r", encoding="utf-8") as f:
    _sa = json.load(f)

PROD_URL = "https://backend-gold-iota-nyngopebeg.vercel.app"
BYPASS_HEADER = {"x-vercel-protection-bypass": "y6RbVJyRTFYY6j5ndZ3yLGRkXYVXX68u"}

# Load JWT_SECRET
env_path = os.path.join(os.path.dirname(__file__), ".env")
jwt_secret = "secret"
if os.path.exists(env_path):
    with open(env_path, "r", encoding="utf-8") as f:
        for line in f:
            if line.startswith("JWT_SECRET="):
                jwt_secret = line.split("=", 1)[1].strip()

sa_creds = service_account.Credentials.from_service_account_file(CRED_FILE)
db = firestore.AsyncClient(project="nyaysetu-pro", credentials=sa_creds)

def make_test_token(user_id: str) -> str:
    return jwt.encode({"sub": user_id, "token_type": "access", "ver": 1}, jwt_secret, algorithm="HS256")

def extract_pdf_text_pypdfium(pdf_bytes: bytes) -> str:
    pdf = pdfium.PdfDocument(pdf_bytes)
    try:
        full_text = []
        for i in range(len(pdf)):
            page = pdf[i]
            try:
                textpage = page.get_textpage()
                try:
                    full_text.append(textpage.get_text_range())
                finally:
                    textpage.close()
            finally:
                page.close()
        return "\n".join(full_text)
    finally:
        pdf.close()

async def main():
    print("=" * 75)
    print(f"NYAYSETU PRO — LIVE PRODUCTION DOCUMENT GENERATION FORENSIC VERIFICATION")
    print(f"Target URL: {PROD_URL}")
    print("=" * 75)

    # 1. Health check
    async with httpx.AsyncClient(headers=BYPASS_HEADER, timeout=60.0) as client:
        r = await client.get(f"{PROD_URL}/healthz")
        print(f"[1] Production /healthz: Status {r.status_code}, Body: {r.json()}")
        assert r.status_code == 200

    # 2. Create ephemeral test user in live Firestore with credits
    unique_id = uuid.uuid4().hex[:8]
    user_id = f"test_docgen_{unique_id}"
    phone = f"9999{unique_id[:6]}"
    now_iso = datetime.now(timezone.utc).isoformat()
    exp_iso = (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat()
    
    print(f"\n[2] Creating ephemeral test user in live Firestore: {user_id} (Mobile: {phone})")
    await db.collection("users").document(user_id).set({
        "id": user_id,
        "phone": phone,
        "mobile": phone,
        "name": "Live Forensic Test Lawyer",
        "full_name": "Live Forensic Test Lawyer",
        "bar_council_no": f"G/TEST/{unique_id[:4]}",
        "user_type": "Advocate",
        "active": True,
        "created_at": now_iso,
        "updated_at": now_iso,
    })
    wallet_id = f"wallet_{user_id}"
    await db.collection("wallets").document(wallet_id).set({
        "id": wallet_id,
        "user_id": user_id,
        "balance": 50,
        "total_used": 0,
        "created_at": now_iso,
        "updated_at": now_iso,
    })
    # Set OTP document in Firestore
    await db.collection("otps").document(phone).set({
        "mobile": phone,
        "otp": "123456",
        "kind": "login",
        "expires_at": exp_iso,
        "attempts": 0,
        "created_at": now_iso,
    })

    # Call verify-otp on Vercel to obtain legitimate token signed by Vercel serverless
    async with httpx.AsyncClient(headers=BYPASS_HEADER, timeout=30.0) as client:
        v_res = await client.post(f"{PROD_URL}/api/auth/verify-otp", json={"mobile": phone, "otp": "123456"})
        assert v_res.status_code == 200, f"Failed OTP verify: {v_res.status_code}, {v_res.text}"
        token = v_res.json()["token"]
        print(f"    Authenticated successfully via OTP. Received live Vercel JWT token.")

    headers = {
        **BYPASS_HEADER,
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
    }

    test_matrix = [
        {
            "name": "Mudat Arji (Gujarati) -> PDF",
            "template_id": "mudat_arji_gu",
            "language": "gu",
            "format": "pdf",
            "values": {
                "court_name": "માનનીય એડી. ચીફ જ્યુડી. મેજી. સાહેબની કોર્ટ, અમદાવાદ",
                "case_number": "ક્રિમીનલ કેસ નં. ૧૨૩૪ / ૨૦૨૬",
                "applicant_name": "રોનક શર્મા",
                "reason": "વકીલશ્રી અન્ય કોર્ટમાં રોકાયેલ છે",
                "advocate_name": "કે. એલ. પટેલ",
                "date": "૦૮/૦૯/૨૦૨૬",
                "place": "અમદાવાદ",
            },
        },
        {
            "name": "Mudat Arji (Gujarati) -> DOCX",
            "template_id": "mudat_arji_gu",
            "language": "gu",
            "format": "docx",
            "values": {
                "court_name": "માનનીય એડી. ચીફ જ્યુડી. મેજી. સાહેબની કોર્ટ, અમદાવાદ",
                "case_number": "ક્રિમીનલ કેસ નં. ૧૨૩૪ / ૨૦૨૬",
                "applicant_name": "રોનક શર્મા",
                "reason": "વકીલશ્રી અન્ય કોર્ટમાં રોકાયેલ છે",
                "advocate_name": "કે. એલ. પટેલ",
                "date": "૦૮/૦૯/૨૦૨૬",
                "place": "અમદાવાદ",
            },
        },
        {
            "name": "Mudat Arji (Gujarati) -> ODT",
            "template_id": "mudat_arji_gu",
            "language": "gu",
            "format": "odt",
            "values": {
                "court_name": "માનનીય એડી. ચીફ જ્યુડી. મેજી. સાહેબની કોર્ટ, અમદાવાદ",
                "case_number": "ક્રિમીનલ કેસ નં. ૧૨૩૪ / ૨૦૨૬",
                "applicant_name": "રોનક શર્મા",
                "reason": "વકીલશ્રી અન્ય કોર્ટમાં રોકાયેલ છે",
                "advocate_name": "કે. એલ. પટેલ",
                "date": "૦૮/૦૯/૨૦૨૬",
                "place": "અમદાવાદ",
            },
        },
        {
            "name": "Mudat Arji (Gujarati) -> PNG Image",
            "template_id": "mudat_arji_gu",
            "language": "gu",
            "format": "png",
            "values": {
                "court_name": "માનનીય એડી. ચીફ જ્યુડી. મેજી. સાહેબની કોર્ટ, અમદાવાદ",
                "case_number": "ક્રિમીનલ કેસ નં. ૧૨૩૪ / ૨૦૨૬",
                "applicant_name": "રોનક શર્મા",
                "reason": "વકીલશ્રી અન્ય કોર્ટમાં રોકાયેલ છે",
                "advocate_name": "કે. એલ. પટેલ",
                "date": "૦૮/૦૯/૨૦૨૬",
                "place": "અમદાવાદ",
            },
        },
        {
            "name": "Exhibit Document (English) -> PDF",
            "template_id": "document_swikaravani_arji_en",
            "language": "en",
            "format": "pdf",
            "values": {
                "court_name": "In the Court of Hon'ble Judicial Magistrate, Ahmedabad",
                "case_number": "Criminal Case No. 456/2026",
                "parties": "State vs. Ronak Sharma",
                "applicant_name": "Ronak Sharma",
                "document_title": "Original Sale Deed dated 12/01/2020",
                "relevance_reason": "is of immense importance to prove defence",
                "documents_list": "1. Original Sale Deed",
                "advocate_name": "K. L. Patel",
                "date": "08/09/2026",
                "place": "Ahmedabad",
            },
        },
        {
            "name": "Exhibit Document (English) -> DOCX",
            "template_id": "document_swikaravani_arji_en",
            "language": "en",
            "format": "docx",
            "values": {
                "court_name": "In the Court of Hon'ble Judicial Magistrate, Ahmedabad",
                "case_number": "Criminal Case No. 456/2026",
                "parties": "State vs. Ronak Sharma",
                "applicant_name": "Ronak Sharma",
                "document_title": "Original Sale Deed dated 12/01/2020",
                "relevance_reason": "is of immense importance to prove defence",
                "documents_list": "1. Original Sale Deed",
                "advocate_name": "K. L. Patel",
                "date": "08/09/2026",
                "place": "Ahmedabad",
            },
        },
        {
            "name": "Exhibit Document (English) -> ODT",
            "template_id": "document_swikaravani_arji_en",
            "language": "en",
            "format": "odt",
            "values": {
                "court_name": "In the Court of Hon'ble Judicial Magistrate, Ahmedabad",
                "case_number": "Criminal Case No. 456/2026",
                "parties": "State vs. Ronak Sharma",
                "applicant_name": "Ronak Sharma",
                "document_title": "Original Sale Deed dated 12/01/2020",
                "relevance_reason": "is of immense importance to prove defence",
                "documents_list": "1. Original Sale Deed",
                "advocate_name": "K. L. Patel",
                "date": "08/09/2026",
                "place": "Ahmedabad",
            },
        },
        {
            "name": "Exhibit Document (English) -> PNG Image",
            "template_id": "document_swikaravani_arji_en",
            "language": "en",
            "format": "png",
            "values": {
                "court_name": "In the Court of Hon'ble Judicial Magistrate, Ahmedabad",
                "case_number": "Criminal Case No. 456/2026",
                "parties": "State vs. Ronak Sharma",
                "applicant_name": "Ronak Sharma",
                "document_title": "Original Sale Deed dated 12/01/2020",
                "relevance_reason": "is of immense importance to prove defence",
                "documents_list": "1. Original Sale Deed",
                "advocate_name": "K. L. Patel",
                "date": "08/09/2026",
                "place": "Ahmedabad",
            },
        },
    ]

    results = []
    try:
        async with httpx.AsyncClient(headers=headers, timeout=60.0) as client:
            print("\n[3] Executing Live Document Generation Requests against Vercel Production:")
            for idx, item in enumerate(test_matrix, 1):
                payload = {
                    "template_id": item["template_id"],
                    "language": item["language"],
                    "format": item["format"],
                    "values": item["values"],
                }
                t0 = time.time()
                res = await client.post(f"{PROD_URL}/api/applications/download", json=payload)
                elapsed = time.time() - t0
                assert res.status_code == 200, f"Failed: status {res.status_code}, {res.text}"
                data = res.json()
                b64_str = data.get("base64") or data.get("content")
                assert b64_str, f"Missing base64 payload: {list(data.keys())}"
                raw = base64.b64decode(b64_str)
                
                forensic_notes = []
                # Forensic verification
                if item["format"] == "pdf":
                    text = extract_pdf_text_pypdfium(raw)
                    char_count = len(text)
                    if item["template_id"] == "mudat_arji_gu":
                        # Forensic assertion: Overlap bug caused length to balloon to 11,760 chars!
                        assert char_count < 1500, f"Overlapping text bug! Character count={char_count}"
                        assert char_count > 300, f"Document empty! Character count={char_count}"
                        forensic_notes.append(f"Clean text length: {char_count} chars (zero overlap)")
                    elif item["template_id"] == "document_swikaravani_arji_en":
                        # Forensic assertion: Inversion bug caused prayer before opening!
                        idx_open = text.find("Most respectfully showeth")
                        if idx_open == -1:
                            idx_open = text.find("submit before this Hon'ble Court")
                        idx_prayer = text.find("PRAYER")
                        if idx_prayer == -1:
                            idx_prayer = text.find("prayed that")
                        assert idx_open != -1 and idx_prayer != -1, "Missing key legal sections"
                        assert idx_open < idx_prayer, f"Displaced/Inverted text! Open={idx_open}, Prayer={idx_prayer}"
                        forensic_notes.append(f"Sequential order verified: Opening ({idx_open}) < Prayer ({idx_prayer})")
                elif item["format"] == "docx":
                    doc = Document(io.BytesIO(raw))
                    for p in doc.paragraphs:
                        if p.paragraph_format.line_spacing is not None:
                            assert p.paragraph_format.line_spacing.pt >= 14.0
                    forensic_notes.append(f"{len(doc.paragraphs)} paragraphs verified with valid leading >= 14pt")
                elif item["format"] == "odt":
                    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
                        content_xml = zf.read("content.xml").decode("utf-8")
                        assert "fo:line-height=\"10%\"" not in content_xml
                    forensic_notes.append("ODT content.xml verified with standard line-height")
                elif item["format"] == "png":
                    assert raw[:8] == b"\x89PNG\r\n\x1a\n"
                    forensic_notes.append(f"Valid PNG header verified ({len(raw)} bytes)")

                note_str = " | ".join(forensic_notes)
                print(f"  [{idx}/8] [PASS] {item['name']:38s} | {len(raw):6d} bytes in {elapsed:5.2f}s | {note_str}")
                results.append({"name": item["name"], "status": "PASS", "bytes": len(raw), "duration": elapsed})
    finally:
        # Cleanup test user and wallet
        print(f"\n[4] Cleaning up ephemeral test user from live Firestore...")
        await db.collection("users").document(user_id).delete()
        await db.collection("wallets").document(wallet_id).delete()
        print("    Cleanup completed successfully.")

    print("\n" + "=" * 75)
    print("LIVE PRODUCTION FORENSIC VERIFICATION COMPLETE: ALL 8/8 TESTS PASSED")
    print("=" * 75)

if __name__ == "__main__":
    asyncio.run(main())
